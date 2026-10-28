import os
import shutil
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Helper: Set Cell Shading
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    # Helper: Add Heading 1
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
        return p

    # Helper: Add Heading 2
    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(59, 130, 246) # Blue 500
        return p

    # Helper: Add Body Text
    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.italic = italic
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # Helper: Add Link Box
    def add_link_item(label, url):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"• {label}: ")
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(30, 41, 59)
        
        r2 = p.add_run(url)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(37, 99, 235)
        r2.font.underline = True

    # ── DOCUMENT HEADER ──────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    t_run = title_p.add_run("RADIX Talent Match — Full Platform Documentation")
    t_run.font.name = "Arial"
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    s_run = sub_p.add_run("End-to-End Candidate Readiness & 12-Skillset Competency Evaluation Platform")
    s_run.font.name = "Arial"
    s_run.font.size = Pt(12)
    s_run.font.color.rgb = RGBColor(100, 116, 139)

    # ── SUMMARY LINKS TABLE ──────────────────────────────────────────────────────
    add_h1("1. Official Project Links & Access")

    link_table = doc.add_table(rows=6, cols=2)
    link_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    link_table.autofit = False

    links_data = [
        ("Live Web Application (Vercel)", "https://radix-talent-match-3ttq.vercel.app"),
        ("Public GitHub Repository", "https://github.com/vatsalawasthi/radix-talent-match"),
        ("Live Backend API Base (Render)", "https://radix-talent-match.onrender.com"),
        ("Interactive Swagger UI (Admin / API Docs)", "https://radix-talent-match.onrender.com/docs"),
        ("ReDoc API Documentation", "https://radix-talent-match.onrender.com/redoc"),
        ("Live MongoDB Telemetry & DB Status", "https://radix-talent-match.onrender.com/api/data/db-status"),
    ]

    for i, (resource, link) in enumerate(links_data):
        row = link_table.rows[i]
        # Col 0: Resource
        c0 = row.cells[0]
        c0.width = Inches(2.6)
        set_cell_background(c0, "F8FAFC" if i % 2 == 0 else "FFFFFF")
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(resource)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(30, 41, 59)

        # Col 1: URL
        c1 = row.cells[1]
        c1.width = Inches(4.2)
        set_cell_background(c1, "F8FAFC" if i % 2 == 0 else "FFFFFF")
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(link)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(37, 99, 235)
        r1.font.underline = True

    # ── EXECUTIVE SUMMARY ────────────────────────────────────────────────────────
    add_h1("2. Executive Summary")
    add_p(
        "RADIX Talent Match is a real, end-to-end full-stack platform built for the RADIX Talent Match Hackathon 2026. "
        "The system flips the lens of traditional hiring tools: instead of evaluating candidates from the employer's viewpoint, "
        "it is designed specifically for the candidate to understand how ready they actually are for specific roles across the "
        "official 12-skillset competency framework."
    )
    add_p(
        "The platform brings together five distinct roles into one seamless, integrated web application: "
        "(1) JD Analytics, (2) Resume Parsing, (3) Profile Builder, (4) Talent Check, and (5) Skill Matching. "
        "It features live cloud deployment on Vercel (Frontend React 18 / Vite), Render (FastAPI / Python 3.11), and "
        "MongoDB Atlas database persistence with zero dummy data, validated against 6 authentic job postings and 4 distinct candidate profiles."
    )

    # ── THE 12 RADIX CATEGORIES ──────────────────────────────────────────────────
    add_h1("3. The 12 RADIX Skill Categories & Shared Data Contract")
    add_p("Every extracted skill is classified with exactly one category_code matching the hackathon standard:")

    cat_table = doc.add_table(rows=13, cols=3)
    cat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cat_table.autofit = False

    cat_data = [
        ("Code", "Category Name", "Description & Key Competencies"),
        ("COD", "Coding", "Programming languages (Python, Java, C++, Go, TypeScript), clean coding, refactoring"),
        ("DSA", "Data Structures & Algorithms", "Space-time complexity, graph traversal, trees, dynamic programming, sorting"),
        ("OOD", "Object-Oriented Design", "SOLID principles, modular architecture, design patterns, domain models"),
        ("APTI", "Aptitude & Problem Solving", "Analytical reasoning, quantitative logic, root-cause analysis (RCA), debugging"),
        ("COMM", "Communication", "Technical documentation, cross-functional collaboration, presentation, client handover"),
        ("AI", "Artificial Intelligence", "Machine learning, PyTorch, deep learning, NLP, transformers, LLMs, embeddings"),
        ("CLOUD", "Cloud Computing", "AWS, GCP, Azure, Docker, Kubernetes, CI/CD, microservices, serverless"),
        ("SQL", "SQL & Databases", "PostgreSQL, Oracle DB, query optimization, stored procedures, indexing, ACID transactions"),
        ("SWE", "Software Engineering", "Unit testing (JUnit, pytest), Git workflows, code reviews, Agile/Scrum ceremonies"),
        ("SYSD", "System Design", "Distributed consensus (Raft), message queues (Kafka), low-latency caching (Redis), scalability"),
        ("NETW", "Networking", "TCP/IP socket programming, HTTP/2 & gRPC, RESTful APIs, DNS, SSL/TLS, eBPF"),
        ("OS", "Operating Systems", "Linux/Unix internals, multithreading, concurrency, process scheduling, POSIX APIs"),
    ]

    for i, (code, name, desc) in enumerate(cat_data):
        row = cat_table.rows[i]
        is_hdr = (i == 0)
        fill = "1E293B" if is_hdr else ("F8FAFC" if i % 2 == 1 else "FFFFFF")

        for j, val in enumerate([code, name, desc]):
            cell = row.cells[j]
            set_cell_background(cell, fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5) if not is_hdr else Pt(10)
            r.font.bold = is_hdr or (j == 0)
            r.font.color.rgb = RGBColor(255, 255, 255) if is_hdr else RGBColor(15, 23, 42)

        row.cells[0].width = Inches(1.0)
        row.cells[1].width = Inches(2.2)
        row.cells[2].width = Inches(3.6)

    # ── THE FIVE CORE MODULES ────────────────────────────────────────────────────
    add_h1("4. Detailed Breakdown of the Five Core Modules")

    add_h2("Module 01: JD Analytics (Role 1 — LLM & NLP Extraction)")
    add_p(
        "• Mission: Ingest any Job Description in PDF, Word (.docx), or plain text format and extract a structured list of technical skills.\n"
        "• Methodology: Concentrates signal extraction on 'Key Responsibilities' and 'What We're Looking For' sections as specified in the participant brief. Uses a 163-parameter semantic mapping dictionary to categorize requirements into the 12 RADIX codes.\n"
        "• Output: Generates the shared JSON schema including category_code, verbatim quote evidence, and confidence rating (high/medium/low).\n"
        "• Integration: Includes a one-click button to push the extracted JD directly into Module 05 for instant role matching."
    )

    add_h2("Module 02: Resume Parsing (Role 2 — LLM & Entity Extraction)")
    add_p(
        "• Mission: Parse candidate resumes across layout variations, pulling both structured credentials and technical skill evidence.\n"
        "• Methodology: Tolerant parser that extracts candidate contact metadata (Name, Email), Education credentials, internship/work experiences, and hackathon project highlights.\n"
        "• Output: Maps detected technical capabilities into the 12 RADIX skill categories with contextual evidence quotes.\n"
        "• Integration: Features a 'Send to Profile Builder' action to automatically pre-fill the candidate profile in Module 03."
    )

    add_h2("Module 03: Profile Builder (Role 3 — Profile & JSON Persistence)")
    add_p(
        "• Mission: Allow candidates to view, edit, and enrich their complete technical profile beyond the raw resume.\n"
        "• Features: Interactive manager for basic details, 12-category skill inventory with dynamic add/remove, hackathon awards, internships, and certifications.\n"
        "• Persistence: Fully persists active profiles locally and in MongoDB (`/data/active_profile.json` and MongoDB collection `candidates`).\n"
        "• Export / Import: Complete JSON import/export matching the Page 10 shared data contract."
    )

    add_h2("Module 04: Talent Check (Role 4 — Company Benchmark & Gap Analysis)")
    add_p(
        "• Mission: Provide a one-button comparison of candidate readiness across all 12 skillsets against specific company hiring bars.\n"
        "• Benchmark Data: Directly consumes `talent_check_company_skillsets.json` containing real competency thresholds (levels 1–10) for Google LLC, Microsoft Corporation, and Oracle Financial Services Software.\n"
        "• Visualizations: Features an interactive 12-category Radar Chart (Candidate vs Company Required) and an animated Readiness Score Ring (0–100).\n"
        "• Gap Breakdown: Details exactly which skillsets meet or clear the bar versus critical gaps requiring preparation."
    )

    add_h2("Module 05: Skill Matching (Role 5 — Role Fuzzy Match & Gap List)")
    add_p(
        "• Mission: Evaluate candidate match against one specific job posting (not a general company bar).\n"
        "• Methodology: Performs fuzzy string similarity, token overlap, and semantic category alignment between candidate evidence and JD requirements.\n"
        "• Output: Generates an overall Match Score (0–100), overall verdict tag (Strong Match, Partial Match, Not Ready), a side-by-side list of Matched Requirements, and a prioritized list of Missing Skills with concrete advice on how to bridge each gap."
    )

    # ── THE INTEGRATED 5-STEP PIPELINE ───────────────────────────────────────────
    add_h1("5. The Integrated 5-Step Pipeline (Page 12)")
    add_p(
        "The platform weaves all five pieces into one continuous guided candidate journey:\n"
        "1. Step 1 (Pick a JD): Select one of the 6 real hackathon job postings or upload a custom JD.\n"
        "2. Step 2 (Parse Resume): Select from the 4 sample candidate resumes or upload a PDF/Word CV.\n"
        "3. Step 3 (Build Profile): Confirm and enrich the pre-filled candidate profile in the shared JSON schema.\n"
        "4. Step 4 (Talent Check): Run one-button company benchmarking to generate the 12-skillset Radar Chart and readiness score.\n"
        "5. Step 5 (Skill Match): Run role-specific fuzzy matching to inspect matched capabilities and prioritized gap advice.\n\n"
        "🚀 1-Click Auto-Demo Runner: The homepage features a 1-Click Auto-Demo button that executes all five stages sequentially in seconds for rapid jury presentations."
    )

    # ── VERIFICATION & AUTHENTIC TEST DATA ────────────────────────────────────────
    add_h1("6. Verification & Authentic Test Data (No Fake Data)")
    add_p(
        "The entire platform was built and tested strictly against the real materials specified in Page 11 of the participant brief. "
        "The scoring logic produces visibly distinct, role-accurate results across all four candidate profiles:"
    )

    test_table = doc.add_table(rows=5, cols=5)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_table.autofit = False

    test_data = [
        ("Candidate", "Profile Archetype (Page 11)", "Google SWE Match", "Google DS Match", "Oracle FSS Match"),
        ("Ananya Rao", "Systems-leaning generalist", "97% (Strong)", "59% (Not Ready)", "92% (Ready)"),
        ("Karthik Subramaniam", "Data-leaning analyst", "51% (Not Ready)", "96% (Strong)", "91% (Ready)"),
        ("Priya Menon", "Entry-level, support-oriented", "66% (Partial)", "63% (Partial)", "95% (Ready)"),
        ("Rohan Verma", "Strong all-rounder", "98% (Strong)", "98% (Strong)", "98% (Ready)"),
    ]

    for i, row_data in enumerate(test_data):
        row = test_table.rows[i]
        is_hdr = (i == 0)
        fill = "1E293B" if is_hdr else ("F8FAFC" if i % 2 == 1 else "FFFFFF")

        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5) if not is_hdr else Pt(10)
            r.font.bold = is_hdr or (j == 0)
            r.font.color.rgb = RGBColor(255, 255, 255) if is_hdr else RGBColor(15, 23, 42)

        row.cells[0].width = Inches(1.6)
        row.cells[1].width = Inches(2.2)
        row.cells[2].width = Inches(1.1)
        row.cells[3].width = Inches(1.1)
        row.cells[4].width = Inches(1.0)

    # ── CLOUD ARCHITECTURE & DEPLOYMENT ──────────────────────────────────────────
    add_h1("7. Cloud Architecture & Tech Stack")
    add_p(
        "• Frontend Hosting: Vercel (Production URL: https://radix-talent-match-3ttq.vercel.app)\n"
        "  - Framework: React 18, Vite 5, Recharts (Radar & Visualizations), TypeScript.\n"
        "  - Routing: Client-side SPA routing with instant API proxy fallback.\n"
        "• Backend API Hosting: Render (Production URL: https://radix-talent-match.onrender.com)\n"
        "  - Framework: FastAPI, Python 3.11, Uvicorn, Pydantic v2.\n"
        "  - Ingestion: PyPDF and python-docx for binary document parsing.\n"
        "• Database Layer: MongoDB (Atlas / Render hosted)\n"
        "  - Driver: pymongo with automatic connection pooling and local JSON resilience.\n"
        "  - Collections: candidates, jds, talent_checks, skill_matches, milestones.\n"
        "• Source Code Repository: GitHub (https://github.com/vatsalawasthi/radix-talent-match)"
    )

    # ── SCRUM MASTER MILESTONES ──────────────────────────────────────────────────
    add_h1("8. Scrum Master Milestone Checklist (Definition of Done)")
    add_p(
        "• Shared Data Contract: Locked and enforced across all 5 pieces.\n"
        "• Module 01 (JD Analytics): Ingests all 6 hackathon JDs and extracts 12-skillset requirements.\n"
        "• Module 02 (Resume Parsing): Ingests all 4 candidate resumes and pulls technical signals.\n"
        "• Module 03 (Profile Builder): Interactive form with save/load persistence in local and MongoDB.\n"
        "• Module 04 (Talent Check): Evaluates candidate readiness against Google, Microsoft, and Oracle benchmarks.\n"
        "• Module 05 (Skill Matching): Generates fuzzy match scores, matched requirements, and gap lists.\n"
        "• End-to-End Integration: All 5 roles wired into the guided candidate flow."
    )

    # Save documents
    out_dir = Path("C:/Users/Vatsal/.gemini/antigravity/scratch/radix-talent-match")
    doc_path1 = out_dir / "RADIX_Talent_Match_Documentation.docx"
    doc.save(str(doc_path1))

    # Also save to Downloads
    downloads_path = Path("C:/Users/Vatsal/Downloads/RADIX_Talent_Match_Documentation.docx")
    try:
        shutil.copy(str(doc_path1), str(downloads_path))
        print("✓ Copied to Downloads:", downloads_path)
    except Exception as e:
        print("Could not copy to downloads:", e)

    # Also copy to frontend public directory so it's downloadable from the live website!
    public_path = out_dir / "frontend" / "public" / "RADIX_Talent_Match_Documentation.docx"
    public_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(str(doc_path1), str(public_path))
    print("✓ Copied to frontend/public for web download:", public_path)

    print("✓ Main Word Document created successfully at:", doc_path1)

if __name__ == "__main__":
    create_document()
